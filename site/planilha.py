"""
site/planilha.py
Gera a pasta de trabalho XLSX publicada junto com o site.

Uso (normalmente chamado por site/build.py):
    python site/planilha.py --saida site/dist/dados/observatorio-psicologia.xlsx

NÚMERO É NÚMERO, AUSÊNCIA É CÉLULA VAZIA
-----------------------------------------
Os valores vão como número, não como texto formatado — quem abre a planilha
precisa poder somar, ordenar e fazer gráfico sem antes limpar o dado. O formato
de exibição vai no `number_format`, que o Excel renderiza na convenção de quem
abre: vírgula decimal aqui, ponto decimal em outro idioma. Um texto "0,547"
gravado à força quebraria as duas coisas.

Ausência é **célula vazia**, nunca zero e nunca o texto "sem dados". No site o
texto é a leitura certa, porque ninguém soma uma página; numa planilha, "sem
dados" numa coluna numérica contamina toda fórmula que a atravesse, e zero
mente. Célula vazia é o que o Excel entende como ausência — `MÉDIA` a ignora,
`SOMA` a ignora, e o gráfico abre um buraco em vez de mergulhar até o eixo.
"""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

import catalogo  # noqa: E402

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("Instale openpyxl: pip install openpyxl")

RAIZ = Path(__file__).resolve().parent.parent

# Mesmas cores institucionais do CSS, no formato ARGB do openpyxl.
INDIGO = "FF262B54"
ARDOSIA = "FF43509A"
PAPEL = "FFF4F5FA"

CABECALHO = Font(name="Calibri", size=11, bold=True, color="FFF3F4FA")
FUNDO_CABECALHO = PatternFill("solid", fgColor=INDIGO)
TITULO = Font(name="Calibri", size=13, bold=True, color=INDIGO)
NOTA = Font(name="Calibri", size=9, italic=True, color="FF5A6172")
BORDA = Border(bottom=Side(style="thin", color="FFE3E3EC"))


def _formato(dec):
    """Formato de exibição a partir das casas decimais declaradas no catálogo."""
    if dec == 0:
        return "#,##0"
    return "#,##0." + "0" * dec


def _escrever_cabecalho(aba, colunas, larguras=None):
    for i, titulo in enumerate(colunas, start=1):
        c = aba.cell(row=1, column=i, value=titulo)
        c.font = CABECALHO
        c.fill = FUNDO_CABECALHO
        c.alignment = Alignment(vertical="center", wrap_text=True)
    aba.row_dimensions[1].height = 30
    # Painel congelado: sem isso, rolar centenas de municípios deixa o leitor sem saber
    # de que coluna é o número que está olhando.
    aba.freeze_panes = "B2"
    aba.auto_filter.ref = None
    for i, largura in enumerate(larguras or [], start=1):
        aba.column_dimensions[get_column_letter(i)].width = largura


def _valor(aba, linha, coluna, valor, dec=None):
    """Grava o valor. None vira célula vazia — nunca 0, nunca texto."""
    if valor is None:
        return
    c = aba.cell(row=linha, column=coluna, value=valor)
    if isinstance(valor, (int, float)) and dec is not None:
        c.number_format = _formato(dec)
    return c


# --------------------------------------------------------------------------- #

def aba_ufs(wb, nacional, nomes_uf):
    aba = wb.create_sheet("Indicadores por UF")
    indicadores = [i for i in catalogo.INDICADORES]
    colunas = ["UF", "Estado", "Região"] + [i["sigla"] for i in indicadores]
    _escrever_cabecalho(aba, colunas, [7, 22, 14] + [15] * len(indicadores))

    for linha, sigla in enumerate(sorted(nacional["ufs"]), start=2):
        d = nacional["ufs"][sigla]
        aba.cell(row=linha, column=1, value=sigla)
        aba.cell(row=linha, column=2, value=nomes_uf.get(sigla, sigla))
        aba.cell(row=linha, column=3, value=d.get("regiao"))
        for j, ind in enumerate(indicadores, start=4):
            valor = d.get(ind["key"])
            if valor is not None and ind.get("mult"):
                valor = valor * ind["mult"]
            _valor(aba, linha, j, valor, ind["dec"])
    aba.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{len(nacional['ufs']) + 1}"
    return aba


def aba_municipios(wb, dir_municipios, nomes_uf):
    aba = wb.create_sheet("Municípios")
    campos = ["vagas_presencial", "cursos_presencial", "n_ies", "matriculas",
              "matriculas_ead", "concluintes", "polos_ead",
              "psicologos_sus", "estabelecimentos_raps", "raps_comunitaria",
              "raps_residencial", "raps_internacao", "cras", "creas"]
    colunas = ["UF", "Estado", "Código IBGE", "Município"] + [
        catalogo.POR_CHAVE[c]["sigla"] for c in campos]
    _escrever_cabecalho(aba, colunas, [7, 20, 14, 30] + [14] * len(campos))

    linha = 2
    for arquivo in sorted(Path(dir_municipios).glob("*.json")):
        sigla = arquivo.stem.upper()
        conteudo = json.loads(arquivo.read_text(encoding="utf-8"))
        for m in conteudo["municipios"]:
            aba.cell(row=linha, column=1, value=sigla)
            aba.cell(row=linha, column=2, value=nomes_uf.get(sigla, sigla))
            # Código como TEXTO: o IBGE usa 7 dígitos e alguns começam com
            # zero à esquerda em outros recortes. Como número, o Excel os come.
            aba.cell(row=linha, column=3, value=str(m["codigo"])).number_format = "@"
            aba.cell(row=linha, column=4, value=m.get("nome"))
            for j, campo in enumerate(campos, start=5):
                _valor(aba, linha, j, m.get(campo), catalogo.POR_CHAVE[campo]["dec"])
            linha += 1
    aba.auto_filter.ref = f"A1:{get_column_letter(len(colunas))}{linha - 1}"
    return aba


def aba_cursos(wb, qualidade):
    if not qualidade:
        return None
    aba = wb.create_sheet("Cursos avaliados (CPC)")
    colunas = ["UF", "Município", "IES", "Sigla", "Categoria", "Modalidade",
               "Cód. curso", "Concluintes part.", "CPC faixa", "CPC contínuo",
               "ENADE contínuo", "IDD", "Vagas"]
    _escrever_cabecalho(aba, colunas, [7, 24, 42, 12, 22, 20, 12, 16, 11, 13, 14, 10, 10])

    for linha, c in enumerate(qualidade["cursos"], start=2):
        aba.cell(row=linha, column=1, value=c.get("uf"))
        aba.cell(row=linha, column=2, value=c.get("municipio"))
        aba.cell(row=linha, column=3, value=c.get("ies"))
        aba.cell(row=linha, column=4, value=c.get("sigla_ies"))
        aba.cell(row=linha, column=5, value=c.get("categoria"))
        aba.cell(row=linha, column=6, value=c.get("modalidade"))
        aba.cell(row=linha, column=7, value=str(c.get("cod_curso"))).number_format = "@"
        _valor(aba, linha, 8, c.get("concluintes_participantes"), 0)
        _valor(aba, linha, 9, c.get("CPC_faixa"), 0)
        _valor(aba, linha, 10, c.get("CPC_cont"), 3)
        _valor(aba, linha, 11, c.get("ENADE_cont"), 3)
        _valor(aba, linha, 12, c.get("IDD"), 3)
        _valor(aba, linha, 13, c.get("vagas"), 0)
    aba.auto_filter.ref = f"A1:M{len(qualidade['cursos']) + 1}"
    return aba


def aba_glossario(wb):
    aba = wb.create_sheet("Glossário")
    colunas = ["Sigla", "Indicador", "Categoria", "O que mede", "Escala",
               "Direção", "Fonte", "Campo no JSON"]
    _escrever_cabecalho(aba, colunas, [16, 34, 15, 78, 18, 14, 22, 26])
    aba.freeze_panes = "A2"

    seta = {"maior": "maior é melhor", "menor": "menor é melhor",
            "contextual": "sem juízo normativo"}
    for linha, i in enumerate(catalogo.INDICADORES, start=2):
        aba.cell(row=linha, column=1, value=i["sigla"]).font = Font(bold=True)
        aba.cell(row=linha, column=2, value=i["nome"])
        aba.cell(row=linha, column=3, value=i["cat"])
        c = aba.cell(row=linha, column=4, value=i["oque"])
        c.alignment = Alignment(wrap_text=True, vertical="top")
        aba.cell(row=linha, column=5, value=i["escala"])
        aba.cell(row=linha, column=6, value=seta.get(i["dir"], i["dir"]))
        aba.cell(row=linha, column=7, value=i["fonte"])
        aba.cell(row=linha, column=8, value=i["key"])
        aba.row_dimensions[linha].height = 46
    return aba


def aba_leia_me(wb, nacional, proveniencia):
    """
    Primeira aba: o que a planilha é, de onde vêm os números e o que falta.

    Uma planilha viaja sem o site. Quem a receber por e-mail não tem o aviso
    legal, nem o glossário na tela, nem a nota sobre o CER — e vai tratar cada
    coluna como fato fechado. As ressalvas precisam viajar junto.
    """
    aba = wb.create_sheet("Leia-me", 0)
    aba.column_dimensions["A"].width = 26
    aba.column_dimensions["B"].width = 96

    meta = nacional["metadados"]
    linhas = [
        ("Observatório Nacional da Formação em Psicologia", None),
        ("", None),
        ("Autoria", "Edson Sidião de Souza Júnior — sidiao@i9educar.com"),
        ("Site", "https://esidiao.github.io/observatorio-psicologia/"),
        ("Ano do Censo", meta.get("ano_censo")),
        ("Gerado em", meta.get("gerado_em")),
        ("", None),
        ("COMO LER", None),
        ("Célula vazia",
         "Ausência de dado — NÃO é zero. Sem fonte oficial para o recorte, "
         "nada é estimado, interpolado ou preenchido por analogia."),
        ("Números",
         "Gravados como número, com formato de exibição. Some, ordene e faça "
         "gráfico à vontade; a vírgula decimal é do seu Excel, não do arquivo."),
        ("Glossário",
         "A aba Glossário define cada indicador: o que mede, em que escala, "
         "se um valor alto é bom, e de que fonte vem."),
        ("", None),
        ("FONTES", None),
    ]

    fontes = (proveniencia or {}).get("fontes", {})
    rotulos = {"censo": "Censo da Educação Superior (INEP)",
               "cpc": "Conceito Preliminar de Curso (INEP)",
               "cnes": "CNES (Ministério da Saúde)",
               "cadsuas": "CadSUAS (Ministério do Desenvolvimento Social)",
               "ibge_municipios": "Municípios (IBGE)",
               "ibge_populacao": "População (IBGE)"}
    for chave, rotulo in rotulos.items():
        f = fontes.get(chave) or {}
        detalhe = (f.get("ano_censo") or f.get("ciclo") or f.get("competencia")
                   or f.get("anomes") or f.get("ano") or "")
        extraido = f.get("extraido_em", "")
        linhas.append((rotulo, f"{detalhe}{'  ·  extraído em ' + extraido if extraido else ''}"))

    linhas += [("", None), ("LIMITAÇÕES CONHECIDAS", None)]
    for lim in (proveniencia or {}).get("limitacoes_conhecidas", []):
        linhas.append(("", lim))

    linhas += [
        ("", None),
        ("DIREITOS", None),
        ("Dados primários",
         "Públicos, do INEP, do Ministério da Saúde, do Ministério do "
         "Desenvolvimento Social e do IBGE."),
        ("Indicadores calculados",
         "Reúso livre, inclusive comercial, citando a fonte primária e este "
         "observatório."),
        ("Obra autoral",
         "Formulação dos índices, catálogo, código e textos: © Edson Sidião de "
         "Souza Júnior, Leis 9.610/1998 e 9.609/1998."),
    ]

    for i, (a, b) in enumerate(linhas, start=1):
        ca = aba.cell(row=i, column=1, value=a)
        if a and a.isupper() and b is None:
            ca.font = Font(bold=True, color=ARDOSIA[2:], size=11)
        elif i == 1:
            ca.font = TITULO
        else:
            ca.font = Font(bold=True) if a else Font()
        if b is not None:
            cb = aba.cell(row=i, column=2, value=b)
            cb.alignment = Alignment(wrap_text=True, vertical="top")
    return aba


def gerar(saida, dir_dados=None, nomes_uf=None):
    dir_dados = Path(dir_dados or (RAIZ / "data"))
    nacional = json.loads((dir_dados / "nacional.json").read_text(encoding="utf-8"))

    def _opcional(nome):
        caminho = dir_dados / nome
        return (json.loads(caminho.read_text(encoding="utf-8"))
                if caminho.exists() else None)

    qualidade = _opcional("qualidade.json")
    proveniencia = _opcional("_proveniencia.json")

    wb = Workbook()
    wb.remove(wb.active)

    aba_leia_me(wb, nacional, proveniencia)
    aba_ufs(wb, nacional, nomes_uf or {})
    aba_municipios(wb, dir_dados / "municipios", nomes_uf or {})
    aba_cursos(wb, qualidade)
    aba_glossario(wb)

    wb.properties.title = "Observatório Nacional da Formação em Psicologia"
    wb.properties.creator = "Edson Sidião de Souza Júnior"
    wb.properties.description = (
        "Indicadores de acesso territorial, qualidade e cobertura assistencial "
        "dos cursos de Psicologia no Brasil.")

    saida = Path(saida)
    saida.parent.mkdir(parents=True, exist_ok=True)
    wb.save(saida)
    return saida


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--saida", default=str(
        RAIZ / "site" / "dist" / "dados" / "observatorio-psicologia.xlsx"))
    p.add_argument("--dados", default=str(RAIZ / "data"))
    args = p.parse_args()

    from build import NOMES_UF
    caminho = gerar(args.saida, args.dados, NOMES_UF)
    print(f"[PLANILHA] -> {caminho} ({caminho.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
